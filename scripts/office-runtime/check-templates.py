"""Independent native-file checks for the generated example catalogue."""
import argparse
import hashlib
import json
import math
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from docx import Document

def check(root, sources):
    catalog = json.loads((root / 'catalog.json').read_text())
    assert {entry.name for entry in root.iterdir()} == {'catalog.json', *(item['id'] for item in catalog)}, 'catalogue resource root'
    numeric_checks, previews = 0, 0
    for item in catalog:
        directory = root / item['id']
        for name, digest in item['files'].items():
            assert hashlib.sha256((directory / name).read_bytes()).hexdigest() == digest, (item['id'], name)
        previews += item['pages']
        sample = directory / item['sampleFile']
        with zipfile.ZipFile(sample) as archive:
            assert archive.testzip() is None
            assert all('\\' not in name and not name.startswith('/') for name in archive.namelist())
            if item['mode'] == 'excel':
                assert all(entry.date_time == (2026, 9, 1, 0, 0, 0) for entry in archive.infolist()), item['id']
                assert b'2026-09-01T00:00:00Z' in archive.read('docProps/core.xml'), item['id']
        if item['mode'] == 'word':
            document = Document(sample)
            assert document.paragraphs, item['id']
            continue
        values, editable = load_workbook(sample, data_only=True), load_workbook(sample, data_only=False)
        for sheet in values:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != 'e', (item['id'], sheet.title, cell.coordinate, cell.value)
        assert sum(len(sheet._charts) for sheet in editable) > 0, item['id']
        legacy = sources / 'legacy' / item['id'] / item['sampleFile']
        if legacy.exists():
            original = load_workbook(legacy, data_only=False)
            assert editable.sheetnames == original.sheetnames
            for sheet in editable:
                assert sheet.freeze_panes == original[sheet.title].freeze_panes
                assert len(sheet._charts) == len(original[sheet.title]._charts)
        else:
            assert all(sheet.freeze_panes == 'A2' for sheet in editable)
        if (directory / 'checks.json').exists():
            for row in json.loads((directory / 'checks.json').read_text()):
                actual = values[row['sheet']][row['cell']].value
                assert math.isclose(actual, row['expected'], abs_tol=row.get('tolerance', 1e-6)), (item['id'], row, actual)
                numeric_checks += 1
    return {'templates': len(catalog), 'preview_pages': previews, 'numeric_checks': numeric_checks,
            'resource_hashes': 'PASS', 'native_zip_crc': 'PASS', 'excel_error_cells': 0,
            'native_charts': 'PASS', 'frozen_panes': 'PASS'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('root', type=Path)
    parser.add_argument('sources', type=Path)
    args = parser.parse_args()
    print(json.dumps(check(args.root, args.sources)))
