"""Build editable examples with the installation's locked Office libraries."""
import argparse
import json
import re
from datetime import datetime
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def meeting_table(document, rows, widths):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            node.set(qn('w:' + key), value)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for column, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width, cell.vertical_alignment = Cm(width), WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', '100'), ('bottom', '100'), ('left', '120'), ('right', '120')]:
                node = OxmlElement('w:' + edge); node.set(qn('w:w'), value); node.set(qn('w:type'), 'dxa'); margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'EAF0F5'); cell._tc.get_or_add_tcPr().append(fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.3
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text); run.bold = index == 0; run.font.size = Pt(10)
    return table


def meeting_word(item, output):
    document = Document()
    document.core_properties.created = document.core_properties.modified = datetime(2026, 9, 1)
    document.core_properties.author = 'DSH Desktop'
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    for name, size, bold in [('Normal', 11, False), ('Title', 23, True), ('Subtitle', 10, False), ('Heading 1', 14, True), ('Heading 2', 11, True)]:
        style = document.styles[name]
        style.font.name, style.font.size, style.font.bold = '微软雅黑', Pt(size), bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(7 if name == 'Normal' else 10)
        if name.startswith('Heading'):
            style.paragraph_format.space_before = Pt(14)
            style.paragraph_format.keep_with_next = True
    for index, page in enumerate(item['pagesContent']):
        if index:
            document.add_page_break()
        document.add_paragraph('会议纪要' if index == 0 else '会议纪要 / 执行跟进', 'Subtitle')
        document.add_paragraph(page['title'], 'Title')
        if page.get('subtitle'):
            document.add_paragraph(page['subtitle'], 'Subtitle')
        for block in page['sections']:
            document.add_paragraph(block['title'], 'Heading 1')
            for text in block.get('text', '').split('\n'):
                if text:
                    document.add_paragraph(text)
            for topic in block.get('topics', []):
                document.add_paragraph(topic['title'], 'Heading 2')
                document.add_paragraph(topic['text'])
            if block.get('table'):
                meeting_table(document, block['table'], block['widths'])
    document.save(output)


def word(item, output):
    if item['id'] == 'meeting-minutes':
        return meeting_word(item, output)
    document = Document()
    document.core_properties.created = document.core_properties.modified = datetime(2026, 9, 1)
    document.core_properties.author = 'DSH Desktop'
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.3)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = '宋体', Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.35
    for name, size in [('Title', 22), ('Heading 1', 16)]:
        style = document.styles[name]
        style.font.name, style.font.size = '黑体', Pt(size)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '黑体')
        style.font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
    for index, page in enumerate(item['pagesContent']):
        if index:
            document.add_page_break()
        document.add_paragraph(page['title'], 'Title')
        if page.get('subtitle'):
            document.add_paragraph(page['subtitle'])
        for block in page['sections']:
            document.add_paragraph(block['title'], 'Heading 1')
            for text in block.get('text', '').split('\n'):
                if text:
                    document.add_paragraph(text)
            if block.get('table'):
                table = document.add_table(rows=0, cols=len(block['table'][0]))
                table.style = 'Table Grid'
                for row in block['table']:
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        cell.text = text
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.save(output)


def normalize_package(path):
    # Kit allocates random chart-axis IDs. Canonical local IDs preserve the
    # axis-reference graph while giving equivalent workbooks stable revisions.
    with ZipFile(path) as archive:
        entries = [(name, archive.read(name)) for name in sorted(archive.namelist())]
    normalized = []
    for name, data in entries:
        if name == 'docProps/core.xml':
            data = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]+', rb'\g<1>2026-09-01T00:00:00Z', data)
        if re.fullmatch(r'xl/charts/chart\d+\.xml', name):
            axis_ids = {}
            def stable_axis(match):
                original = match.group(2)
                axis_ids.setdefault(original, str(100000 + len(axis_ids)).encode())
                return match.group(1) + axis_ids[original] + match.group(3)
            data = re.sub(rb'(<(?:[\w.-]+:)?(?:axId|crossAx)\b[^>]*\bval=")(\d+)(")', stable_axis, data)
        normalized.append((name, data))
    with ZipFile(path, 'w', compression=ZIP_DEFLATED) as archive:
        for name, data in normalized:
            info = ZipInfo(name, (2026, 9, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)


def main():
    import sys
    if len(sys.argv) == 3 and sys.argv[1] == '--normalize-package':
        normalize_package(Path(sys.argv[2]))
        return
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('output')
    args = parser.parse_args()
    root = Path(args.output)
    for item in json.loads(Path(args.source).read_text()):
        directory = root / item['id']
        directory.mkdir(parents=True, exist_ok=True)
        assert item['mode'] == 'word', 'Common examples use the Word generator'
        sample = directory / 'example.docx'
        word(item, sample)
        normalize_package(sample)
        guidance = '检查章节、决议及行动项。保持可编辑文字和表格。'
        if item['id'] == 'meeting-minutes':
            guidance += ('\n\n沿用 A4 纵向、约 2cm 页边距、23pt 黑色主标题、14pt 一级标题与 11pt 正文。'
                         '行动项集中为一张表，采用浅色表头、细灰边框、自动行高、重复表头及约 10pt 表格文字。'
                         '任务与交付标准左对齐，负责人和期限居中。保留案例的标题与表格样式，按内容自然分页。'
                         '\n\n已定事项、讨论建议与待确认事项分别呈现。责任人、期限及条件以用户原文为依据，缺失信息使用待确认。'
                         '完成后用官方 Kit 导出 PDF 并逐页检查，交付可编辑 DOCX。')
        (directory / 'design.md').write_text(
            f"# {item['title']}\n\n{item['description']}\n\n保留当前案例的版式和结构，以用户当前材料替换全部示例事实。未知字段保留为待确认。\n\n" + guidance + '\n')



if __name__ == '__main__':
    main()
